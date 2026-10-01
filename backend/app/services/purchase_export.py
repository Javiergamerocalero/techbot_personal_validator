"""Exportación a Excel del listado de compras de empleados.

Mismo criterio que la plantilla de empleados (openpyxl), para que el
admin descargue lo que ve en pantalla, ya filtrado por fechas y, si
corresponde, por empleado. Pedido por Javier el 2026-10-01.
"""

from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Zona Lima para mostrar las fechas como las espera el operador.
_LIMA_TZ = timezone(timedelta(hours=-5), name="America/Lima")

_HEADERS = [
    ("Fecha", 20),
    ("Empleado", 34),
    ("Código", 16),
    ("Documento", 16),
    ("Monto (S/)", 14),
    ("Kiosco", 22),
]


def build_purchases_xlsx(
    items: list[dict],
    total_cents: int,
    *,
    date_from: date | None,
    date_to: date | None,
    employee_name: str | None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Compras"

    azul = PatternFill("solid", fgColor="1F3A5F")
    blanco = Font(color="FFFFFF", bold=True)

    # Encabezado informativo del filtro aplicado.
    rango = _fmt_rango(date_from, date_to)
    ws.append([f"Compras de empleados — {rango}"])
    ws.append([f"Filtro: {employee_name or 'Todos los empleados'}"])
    ws.append([])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_HEADERS))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(_HEADERS))
    ws["A1"].font = Font(bold=True, size=13)

    header_row = 4
    for col, (title, width) in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.fill = azul
        cell.font = blanco
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    for item in items:
        ws.append(
            [
                _fmt_dt(item["purchased_at"]),
                item["employee_name"],
                item["employee_code"],
                item["document_number"],
                round(item["amount_cents"] / 100, 2),
                item.get("kiosk_name") or "",
            ]
        )

    # Fila de total.
    total_row = header_row + 1 + len(items) + 1
    ws.cell(row=total_row, column=4, value="TOTAL").font = Font(bold=True)
    tot = ws.cell(row=total_row, column=5, value=round(total_cents / 100, 2))
    tot.font = Font(bold=True)
    ws.cell(row=total_row, column=2, value=f"{len(items)} registro(s)")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt_dt(dt: datetime) -> str:
    local = dt.astimezone(_LIMA_TZ) if dt.tzinfo else dt.replace(tzinfo=timezone.utc).astimezone(_LIMA_TZ)
    return local.strftime("%d/%m/%Y %H:%M")


def _fmt_rango(date_from: date | None, date_to: date | None) -> str:
    if date_from and date_to:
        return f"{date_from.strftime('%d/%m/%Y')} al {date_to.strftime('%d/%m/%Y')}"
    if date_from:
        return f"desde {date_from.strftime('%d/%m/%Y')}"
    if date_to:
        return f"hasta {date_to.strftime('%d/%m/%Y')}"
    return "todas las fechas"
